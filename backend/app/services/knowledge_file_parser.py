"""知识库上传文件解析服务。

第一版支持：
- Markdown / TXT：直接按文本解析，自动尝试 UTF-8 / GB18030。
- PDF：使用 pypdf 提取可复制文本。
- Word docx：使用 python-docx 提取段落和表格文本。

设计取舍：
- 这里不做 OCR；扫描版 PDF 后续可以接 PaddleOCR / 云 OCR。
- 解析服务只负责把文件变成纯文本，后续 chunk、embedding、入库仍交给 KnowledgeService。
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any


MAX_UPLOAD_BYTES = 20 * 1024 * 1024


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".pdf",
    ".docx",
}


@dataclass(frozen=True)
class ParsedKnowledgeFile:
    """上传文件解析结果。"""

    title: str
    text: str
    file_name: str
    file_ext: str
    content_type: str | None
    metadata: dict[str, Any]


def parse_upload_tags(tags: str | list[str] | None) -> list[str]:
    """把表单中的 tags 解析为列表。

    支持：
    - "AI,产品情报,技术趋势"
    - "AI 产品情报 技术趋势"
    - ["AI", "开源"]
    """
    if tags is None:
        return []
    if isinstance(tags, list):
        return [item.strip() for item in tags if str(item).strip()]
    normalized = str(tags).replace("，", ",").replace("；", ",").replace(";", ",")
    result: list[str] = []
    for part in normalized.replace("\n", ",").split(","):
        for item in part.split():
            item = item.strip()
            if item:
                result.append(item)
    return result


def parse_knowledge_file(file_name: str, content: bytes, content_type: str | None = None) -> ParsedKnowledgeFile:
    """解析上传文件为纯文本。"""
    safe_name, ext = validate_knowledge_file_input(file_name, content)

    if ext in {".txt", ".md", ".markdown"}:
        text = _decode_text(content)
    elif ext == ".pdf":
        text = _extract_pdf_text(content)
    elif ext == ".docx":
        text = _extract_docx_text(content)
    else:  # pragma: no cover
        raise ValueError(f"暂不支持的文件类型：{ext}")

    text = _normalize_text(text)
    if len(text) < 10:
        raise ValueError("文件解析后文本太短，可能是扫描版 PDF 或空文档")

    return ParsedKnowledgeFile(
        title=Path(safe_name).stem or "未命名文档",
        text=text,
        file_name=safe_name,
        file_ext=ext.lstrip("."),
        content_type=content_type,
        metadata={
            "file_name": safe_name,
            "file_ext": ext.lstrip("."),
            "content_type": content_type,
            "size_bytes": len(content),
            "parser": "knowledge_file_parser_v1",
        },
    )


def validate_knowledge_file_input(file_name: str, content: bytes) -> tuple[str, str]:
    """Validate upload name and size without doing document parsing."""
    safe_name = Path(file_name or "untitled").name
    ext = Path(safe_name).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"暂不支持的文件类型：{ext or 'unknown'}；当前支持 txt、md、pdf、docx")
    if not content:
        raise ValueError("上传文件为空")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError(f"文件过大：最大支持 {MAX_UPLOAD_BYTES // 1024 // 1024}MB")
    return safe_name, ext


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except Exception as exc:  # noqa: BLE001
        raise ValueError("缺少 pypdf 依赖，无法解析 PDF") from exc

    reader = PdfReader(BytesIO(content))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(f"## 第 {index} 页\n{page_text}")
    return "\n\n".join(pages)


def _extract_docx_text(content: bytes) -> str:
    try:
        from docx import Document
    except Exception as exc:  # noqa: BLE001
        raise ValueError("缺少 python-docx 依赖，无法解析 Word docx") from exc

    doc = Document(BytesIO(content))
    blocks: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _normalize_text(text: str) -> str:
    lines = [line.strip() for line in (text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    cleaned: list[str] = []
    blank_seen = False
    for line in lines:
        if not line:
            if not blank_seen:
                cleaned.append("")
            blank_seen = True
            continue
        cleaned.append(line)
        blank_seen = False
    return "\n".join(cleaned).strip()
